"""Verify generated Excel workbooks"""
import sys
sys.path.insert(0, '.')
from openpyxl import load_workbook

print("=" * 60)
print("EXCEL WORKBOOK VERIFICATION")
print("=" * 60)

# Verify GRANIT workbook
print("\n1. GRANIT_Parts_EU_Sample_Dataset.xlsx")
try:
    wb = load_workbook('output/GRANIT_Parts_EU_Sample_Dataset.xlsx')
    print("   File exists and opens successfully")
    print(f"   Sheet names: {wb.sheetnames}")
    print(f"   Total sheets: {len(wb.sheetnames)}")
    
    expected_granth_sheets = [
        "Products", "Regional Pricing", "Specifications",
        "Compliance", "Inventory & Logistics", "Data Quality",
        "Data Dictionary", "Sources & Assumptions"
    ]
    
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"   - {sheet}: {ws.max_row} rows x {ws.max_column} cols")
        
        # Check first sheet (Products) for headers
        if sheet == "Products":
            headers = []
            for c in range(1, ws.max_column + 1):
                headers.append(ws.cell(row=1, column=c).value)
            print(f"     Headers: {headers}")
            
except Exception as e:
    print(f"   Error: {type(e).__name__}: {e}")
    import traceback
    traceback.print_exc()

# Verify Blinkit workbook
print("\n2. Blinkit_Bengaluru_Sample_Dataset.xlsx")
try:
    wb = load_workbook('output/Blinkit_Bengaluru_Sample_Dataset.xlsx')
    print("   File exists and opens successfully")
    print(f"   Sheet names: {wb.sheetnames}")
    print(f"   Total sheets: {len(wb.sheetnames)}")
    
    expected_blinkit_sheets = [
        "Products", "Pricing & Tax", "Packaging & Ingredients",
        "Suppliers", "Dark Stores & Inventory", "Delivery & Zones",
        "Data Quality", "Data Dictionary", "Sources & Assumptions"
    ]
    
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"   - {sheet}: {ws.max_row} rows x {ws.max_column} cols")
        
        # Check first sheet (Products) for headers
        if sheet == "Products":
            headers = []
            for c in range(1, ws.max_column + 1):
                headers.append(ws.cell(row=1, column=c).value)
            print(f"     Headers: {headers}")
            
except Exception as e:
    print(f"   Error: {type(e).__name__}: {e}")
    import traceback
    traceback.print_exc()

print("\n" + "=" * 60)
print("EXCEL VERIFICATION COMPLETE")
print("=" * 60)